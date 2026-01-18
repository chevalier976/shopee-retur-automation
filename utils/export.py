"""Excel export module for Shopee return data."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
import io


def export_to_excel(df: pd.DataFrame) -> bytes:
    """
    Export DataFrame to Excel with formatting.
    
    Args:
        df: DataFrame to export
        
    Returns:
        bytes: Excel file content as bytes
    """
    # Create a new workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Retur Shopee"
    
    # Write dataframe to worksheet
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            
            # Format header row
            if r_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        
        for cell in column:
            try:
                if cell.value:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            except:
                pass
        
        # Set column width (with some padding and max limit)
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to bytes
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    return excel_buffer.getvalue()


def generate_filename() -> str:
    """
    Generate filename with timestamp.
    
    Returns:
        str: Filename in format Retur_Shopee_YYYYMMDD_HHMMSS.xlsx
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"Retur_Shopee_{timestamp}.xlsx"
